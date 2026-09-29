"""Fig. 1 (revised): scale-persistence map + held-state concept + CAP design handle.

Data for panel a come from data/coacervate_scale_dataset_DRAFT.xlsx (sheet 'Dataset')
and data/fig1a_envelopes.csv (grouping of dataset rows into envelopes, plus placeholders
for entries not yet in the dataset). Everything else is schematic.
Output: Fig1.svg, Fig1.pdf, Fig1_preview.png (183 mm wide, built 1:1).
"""
import csv
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Rectangle, Wedge

HERE = Path(__file__).resolve().parent
MM = 1 / 25.4
mpl.rcParams.update({
    "font.family": "sans-serif", "font.sans-serif": ["Arial", "Liberation Sans", "DejaVu Sans"],
    "font.size": 6, "axes.labelsize": 7, "xtick.labelsize": 6, "ytick.labelsize": 6,
    "axes.linewidth": 0.5, "lines.linewidth": 1,
    "xtick.major.width": 0.5, "ytick.major.width": 0.5, "xtick.major.size": 2, "ytick.major.size": 2,
    "xtick.direction": "in", "ytick.direction": "in", "xtick.top": True, "ytick.right": True,
    "mathtext.default": "regular",
    "pdf.fonttype": 42, "svg.fonttype": "none", "savefig.dpi": 600,
})

# Okabe-Ito for data encodings
BLUE, SKY, VERM, PURPLE, GREY, BLACK = "#0072B2", "#56B4E9", "#D55E00", "#CC79A7", "#7F7F7F", "#000000"
# Illustration palette (matches the application icon set)
T_FILL, T_EDGE, WATER = "#9FDBD1", "#2A9D8F", "#EAF4FB"
I_ORANGE, I_NAVY, I_YELLOW, I_GREY, I_BROWN = "#E76F51", "#3D5A80", "#F4B942", "#C5CAD6", "#6B4F3A"

FIG_W, FIG_H = 183, 114


def fig_mm(w, h):
    return plt.figure(figsize=(w * MM, h * MM))


def ax_mm(fig, x, y, w, h):
    W, H = fig.get_size_inches() / MM
    return fig.add_axes([x / W, 1 - (y + h) / H, w / W, h / H])


def letter(fig, x, y, s):
    W, H = fig.get_size_inches() / MM
    fig.text(x / W, 1 - y / H, s, fontsize=8, fontweight="bold", ha="left", va="top")


# ---------------------------------------------------------------- data
def load_envelopes():
    wb = openpyxl.load_workbook(HERE / "data" / "coacervate_scale_dataset_DRAFT.xlsx", data_only=True)
    ws = wb["Dataset"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    col = {h: i for i, h in enumerate(hdr)}
    ds = {}
    for r in rows[1:]:
        if r[0] and isinstance(r[col["Batch volume low (L)"]], (int, float)):
            ds[r[0]] = (r[col["Batch volume low (L)"]], r[col["Batch volume high (L)"]],
                        r[col["Liquid-state persistence low (s)"]], r[col["Liquid-state persistence high (s)"]])
    env = {}
    with open(HERE / "data" / "fig1a_envelopes.csv", newline="") as f:
        for e in csv.DictReader(f):
            ids = [i for i in e["from_dataset_ids"].split(";") if i]
            if ids:
                v = np.array([ds[i] for i in ids], dtype=float)
                lo_v, hi_v, lo_t, hi_t = v[:, 0].min(), v[:, 1].max(), v[:, 2].min(), v[:, 3].max()
            else:
                lo_v, hi_v, lo_t, hi_t = (float(e[k]) for k in ("vol_lo_L", "vol_hi_L", "life_lo_s", "life_hi_s"))
            e["box"] = np.log10([lo_v, hi_v, lo_t, hi_t])
            env[e["key"]] = e
    return env


# ---------------------------------------------------------------- icons (mm coordinates, y down)
def drop(ax, cx, cy, r, fc=T_FILL, ec=T_EDGE, lw=0.6, z=3):
    ax.add_patch(Circle((cx, cy), r, fc=fc, ec=ec, lw=lw, zorder=z))


def dot(ax, cx, cy, r, c, z=4, ec="none", lw=0):
    ax.add_patch(Circle((cx, cy), r, fc=c, ec=ec, lw=lw, zorder=z))


def icon_extraction(ax, cx, cy, s):
    r = s * 0.30
    dx = s * 0.12
    drop(ax, cx + dx, cy, r)
    for ox, oy in [(-0.1, -0.12), (0.12, -0.08), (-0.05, 0.12), (0.13, 0.13)]:
        dot(ax, cx + dx + ox * s, cy + oy * s, s * 0.055, I_ORANGE)
    dot(ax, cx - s * 0.42, cy, s * 0.06, I_ORANGE)
    ax.annotate("", xy=(cx + dx - r - s * 0.02, cy), xytext=(cx - s * 0.33, cy),
                arrowprops=dict(arrowstyle="-|>", lw=0.6, color=I_NAVY, mutation_scale=3), zorder=5)
    for oy in (-0.3, 0.3):
        ax.add_patch(Rectangle((cx - s * 0.47, cy + oy * s - s * 0.05), s * 0.1, s * 0.1, fc=I_GREY, ec="none"))


def icon_microreactor(ax, cx, cy, s):
    r = s * 0.33
    drop(ax, cx, cy, r)
    ax.add_patch(Wedge((cx - s * 0.05, cy), s * 0.14, 35, 325, fc=I_NAVY, ec="none", zorder=4))
    dot(ax, cx + s * 0.14, cy, s * 0.05, I_ORANGE)
    dot(ax, cx - s * 0.45, cy, s * 0.05, I_ORANGE)
    d = s * 0.07
    ax.add_patch(plt.Polygon([(cx + s * 0.45, cy - d), (cx + s * 0.45 + d, cy), (cx + s * 0.45, cy + d),
                              (cx + s * 0.45 - d, cy)], fc=I_YELLOW, ec=I_NAVY, lw=0.4, zorder=4))


def icon_sensing(ax, cx, cy, s):
    ax.add_patch(Circle((cx, cy), s * 0.47, fc=I_ORANGE, alpha=0.12, ec="none"))
    ax.add_patch(Circle((cx, cy), s * 0.39, fc=I_ORANGE, alpha=0.18, ec="none"))
    drop(ax, cx, cy, s * 0.30)
    for ox, oy in [(-0.09, -0.08), (0.09, -0.09), (-0.08, 0.09), (0.1, 0.08)]:
        dot(ax, cx + ox * s, cy + oy * s, s * 0.05, I_ORANGE)


def icon_delivery(ax, cx, cy, s):
    drop(ax, cx - s * 0.04, cy, s * 0.29)
    t = np.linspace(-1, 1, 30)
    ax.plot(cx - s * 0.04 + t * s * 0.15, cy - s * 0.04 + np.sin(t * 3 * np.pi) * s * 0.035, color=I_NAVY, lw=0.6, zorder=4)
    dot(ax, cx - s * 0.02, cy + s * 0.12, s * 0.045, I_ORANGE)
    th = np.linspace(-0.55, 0.55, 20)
    ax.plot(cx + s * 0.62 - np.cos(th) * s * 0.28, cy + np.sin(th) * s * 0.45, color=I_NAVY, lw=0.9, zorder=4)
    for oy in (-0.1, 0, 0.1):
        ax.plot([cx - s * 0.5, cx - s * 0.4], [cy + oy * s] * 2, color=I_NAVY, lw=0.5)


def icon_protocell(ax, cx, cy, s):
    for a in np.linspace(0, 2 * np.pi, 28, endpoint=False):
        dot(ax, cx + np.cos(a) * s * 0.44, cy + np.sin(a) * s * 0.44, s * 0.025, I_NAVY)
    drop(ax, cx, cy, s * 0.34)
    dot(ax, cx - s * 0.08, cy - s * 0.08, s * 0.1, I_YELLOW, ec=I_ORANGE, lw=0.4)
    t = np.linspace(0, 1, 20)
    ax.plot(cx + (t - 0.3) * s * 0.22, cy + s * 0.12 + np.sin(t * 2 * np.pi) * s * 0.03, color=I_ORANGE, lw=0.5, zorder=4)


def icon_adhesive(ax, cx, cy, s):
    w, h = s * 0.8, s * 0.2
    for oy in (-0.17, 0.17):
        ax.add_patch(FancyBboxPatch((cx - w / 2, cy + oy * s - h / 2), w, h, boxstyle="round,pad=0,rounding_size=0.3",
                                    fc=I_GREY, ec=I_NAVY, lw=0.5, zorder=4))
    ax.add_patch(Ellipse((cx, cy), w * 0.95, s * 0.16, fc=T_FILL, ec=T_EDGE, lw=0.5, zorder=3))


def icon_hair(ax, cx, cy, s):
    t = np.linspace(-0.45, 0.45, 40)
    for oy in (-0.15, 0.17):
        ax.plot(cx + t * s, cy + oy * s + np.sin(t * 7) * s * 0.05, color=I_BROWN, lw=1.1, solid_capstyle="round", zorder=3)
        for ox in (-0.2, 0.18):
            ax.add_patch(Ellipse((cx + ox * s, cy + oy * s + np.sin(ox * 7) * s * 0.05 - s * 0.03), s * 0.16, s * 0.1,
                                 fc=T_FILL, ec=T_EDGE, lw=0.4, zorder=4))


def icon_capsule(ax, cx, cy, s):
    drop(ax, cx - s * 0.08, cy - s * 0.05, s * 0.3)
    dot(ax, cx - s * 0.08, cy - s * 0.05, s * 0.2, I_YELLOW, ec="#B7791F", lw=0.4)
    drop(ax, cx + s * 0.3, cy + s * 0.26, s * 0.14)
    dot(ax, cx + s * 0.3, cy + s * 0.26, s * 0.08, I_YELLOW, ec="#B7791F", lw=0.4)


# ---------------------------------------------------------------- figure
fig = fig_mm(FIG_W, FIG_H)
ov = ax_mm(fig, 0, 0, FIG_W, FIG_H)  # overlay in mm, y down
ov.set_xlim(0, FIG_W); ov.set_ylim(FIG_H, 0); ov.axis("off"); ov.set_zorder(10)

# ======================= panel a: scale-persistence map
AX, AY, AW, AH = 14, 25, 96, 72
ax = ax_mm(fig, AX, AY, AW, AH)
XL, YL = (-7, 5), (-0.2, 7.6)
ax.set_xlim(*XL); ax.set_ylim(*YL)
ax.set_xticks(range(-6, 5, 2)); ax.set_xticklabels([f"$10^{{{k}}}$" if k else "1" for k in range(-6, 5, 2)])
yt = {"1 s": 0, "1 min": np.log10(60), "1 h": np.log10(3600), "1 d": np.log10(86400),
      "1 mo": np.log10(2.63e6), "1 yr": np.log10(3.156e7)}
ax.set_yticks(list(yt.values())); ax.set_yticklabels(list(yt.keys()))
ax.set_xlabel("Batch volume (L)", labelpad=2)
ax.set_ylabel("Lifetime of the liquid dispersion", labelpad=2)
ax.tick_params(top=True, right=True, pad=1.5)


def d2mm(x, y):
    """data coords of panel a -> overlay mm."""
    return AX + (x - XL[0]) / (XL[1] - XL[0]) * AW, AY + AH - (y - YL[0]) / (YL[1] - YL[0]) * AH


# target region: litre scale and held >= 1 d
x1L, y1d = 0.0, np.log10(86400)
ax.add_patch(Rectangle((x1L, y1d), XL[1] - x1L, YL[1] - y1d, fc=VERM, alpha=0.06, ec="none", zorder=0))
ax.plot([x1L, x1L], [y1d, YL[1]], ls=(0, (2, 1.5)), lw=0.5, color=VERM, zorder=1)
ax.plot([x1L, XL[1]], [y1d, y1d], ls=(0, (2, 1.5)), lw=0.5, color=VERM, zorder=1)
ax.text(XL[1] - 0.15, y1d + 0.1, "Litre scale, held for days", ha="right", va="bottom", fontsize=6, color=VERM)

env = load_envelopes()
STYLE = {  # fill, edge, linestyle, filled?
    "bare":       dict(fc=SKY, alpha=0.30, ec=BLUE, ls="-"),
    "membranized": dict(fc=BLUE, alpha=0.30, ec=BLUE, ls="-"),
    "adhesive":   dict(fc="none", alpha=1, ec=PURPLE, ls="-"),
    "bioprocess": dict(fc="none", alpha=1, ec=GREY, ls=(0, (2.5, 1.5))),
    "capsule":    dict(fc="none", alpha=1, ec=BLACK, ls="-"),
    "haircare":   dict(fc="none", alpha=1, ec=BLACK, ls="-"),
    "thiswork":   dict(fc=VERM, alpha=0.30, ec=VERM, ls="-"),
}
PAD = 1.05  # envelope drawn slightly larger than the range box; min size for near-point entries
for k, st in STYLE.items():
    x0, x1, y0, y1 = env[k]["box"]
    w, h = max((x1 - x0) * PAD, 0.9), max((y1 - y0) * PAD, 0.55)
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    if st["fc"] != "none":
        ax.add_patch(Ellipse((cx, cy), w, h, fc=st["fc"], alpha=st["alpha"], ec="none", zorder=2))
    ax.add_patch(Ellipse((cx, cy), w, h, fc="none", ec=st["ec"], lw=0.75, ls=st["ls"], zorder=3))

# this work: diamond + range bars
x0, x1, y0, y1 = env["thiswork"]["box"]
cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
ax.plot([x0, x1], [cy, cy], color=VERM, lw=0.75, zorder=5)
ax.plot([cx, cx], [y0, y1], color=VERM, lw=0.75, zorder=5)
ax.plot(cx, cy, marker="D", ms=4, mfc=VERM, mec="white", mew=0.4, zorder=6)
ax.text(cx, y0 - 0.12, "This work", ha="center", va="top", fontsize=6, fontweight="bold", color=VERM)
# previous CAP paper
nx, _, ny, _ = env["natchem"]["box"]
ax.plot(nx, ny, marker="o", ms=3.5, mfc="white", mec=VERM, mew=0.75, zorder=6)
ax.text(nx + 0.25, ny + 0.62, "CAP, ref. (2025)", ha="left", va="bottom", fontsize=5.5, color=VERM)
ax.add_patch(FancyArrowPatch((nx + 0.3, ny + 0.05), (cx - 0.55, cy - 0.02), arrowstyle="-|>", mutation_scale=5,
                             lw=0.75, color=VERM, connectionstyle="arc3,rad=-0.12", zorder=5))

# labels (data coords)
lab = dict(fontsize=6, ha="center", va="center", zorder=7)
ax.text(-4.3, 2.6, "Bare\ncoacervates", color=BLUE, **lab)
ax.text(-3.75, 5.45, "Membranized", color="white", fontweight="bold", **lab)
ax.text(-0.3, 0.6, "ELP, ATPS*", color=GREY, **lab)
ax.text(-0.3, 4.05, "Wet adhesives", color=PURPLE, **lab)
ax.text(3.0, 3.95, "Microencapsulation", color=BLACK, **lab)
ax.text(2.85, 0.9, "Hair care", color=BLACK, fontsize=6, ha="right", va="center", zorder=7)

# in-plot icons (overlay mm) for pre-industrial and industrial entries
for fn, (x, y) in [(icon_adhesive, (-1.62, 4.05)), (icon_capsule, (4.55, 3.2)), (icon_hair, (3.74, 0.9))]:
    fn(ov, *d2mm(x, y), 5.2)

# encoding key, below panel a
ky = AY + AH + 10.8
ov.add_patch(Ellipse((AX + 2, ky), 3.2, 1.8, fc=SKY, alpha=0.35, ec="none"))
ov.add_patch(Ellipse((AX + 2, ky), 3.2, 1.8, fc="none", ec=BLUE, lw=0.6))
ov.text(AX + 4.3, ky, "Held dispersion required", va="center", fontsize=5.5)
ov.add_patch(Ellipse((AX + 32, ky), 3.2, 1.8, fc="none", ec=BLACK, lw=0.6))
ov.text(AX + 34.3, ky, "Not required (formed in use, crosslinked, bulk or settled)", va="center", fontsize=5.5)
ky2 = ky + 3.0
for x, c, t, ls in [(AX, BLUE, "Emerging", "-"), (AX + 16, PURPLE, "Pre-industrial", "-"), (AX + 36, BLACK, "Industrial", "-"),
                    (AX + 52, GREY, "Bioprocess", (0, (2.5, 1.5))), (AX + 69, VERM, "This work", "-")]:
    ov.plot([x + 0.4, x + 3.6], [ky2, ky2], color=c, lw=0.9, ls=ls)
    ov.text(x + 4.3, ky2, t, va="center", fontsize=5.5, color=c)

# icon strip above panel a: the five emerging uses
names = ["Extraction", "Microreactors", "Sensing", "Drug delivery", "Protocells"]
fns = [icon_extraction, icon_microreactor, icon_sensing, icon_delivery, icon_protocell]
sx0, step, sy = AX + 7, 12.5, 12.0
for i, (n, fn) in enumerate(zip(names, fns)):
    fn(ov, sx0 + i * step, sy, 7.0)
    ov.text(sx0 + i * step, sy + 5.0, n, ha="center", va="top", fontsize=5.5)
bx0, bx1 = sx0 - 5, sx0 + 4 * step + 5
ov.plot([bx0, bx0, bx1, bx1], [5.3, 4.5, 4.5, 5.3], color=BLUE, lw=0.5)
ov.text((bx0 + bx1) / 2, 4.0, "Emerging uses: need a held dispersion, made at µL–mL", ha="center", va="bottom",
        fontsize=5.5, color=BLUE)
ov.text(AX + AW, AY + AH + 6.6, "*segregative (aqueous two-phase), shown for Fig. 6",
        ha="right", va="top", fontsize=5, color=GREY)

# ======================= panel b: waypoint vs destination
BX, BY = 119, 4  # origin of panel b in mm


def vial(x, y, w=9, h=12, kind="solution", seed=0):
    rng = np.random.default_rng(seed)
    ov.add_patch(Rectangle((x, y), w, h, fc=WATER, ec="#555555", lw=0.5, zorder=2))
    ov.add_patch(Rectangle((x - 0.4, y - 1.2), w + 0.8, 1.2, fc="#555555", ec="none", zorder=3))
    if kind == "solution":
        for _ in range(26):
            dot(ov, x + rng.uniform(0.8, w - 0.8), y + rng.uniform(1, h - 0.8), 0.16, T_EDGE)
    elif kind == "transient":
        for r, px, py in [(1.3, 0.3, 0.65), (0.6, 0.7, 0.35), (0.9, 0.3, 0.3), (0.45, 0.72, 0.72),
                          (1.6, 0.62, 0.86), (0.4, 0.25, 0.9), (0.55, 0.8, 0.12)]:
            drop(ov, x + px * w, y + py * h, r, lw=0.4)
    elif kind == "bulk":
        ov.add_patch(Rectangle((x, y + h * 0.62), w, h * 0.38, fc=T_FILL, ec="none", zorder=2.5))
        ov.plot([x, x + w], [y + h * 0.62] * 2, color=T_EDGE, lw=0.5, zorder=3)
    elif kind == "held":
        for i in range(4):
            for j in range(5):
                drop(ov, x + 1.4 + i * 2.05 + (j % 2) * 0.5, y + 1.6 + j * 2.25, 0.72, ec=VERM, lw=0.6)


ov.text(BX, BY + 1, "No stabilizer", fontsize=6, color=GREY, va="center")
vx = [BX + 2, BX + 26, BX + 50]
vy = BY + 5
for x, k, t in zip(vx, ["solution", "transient", "bulk"], ["Solution", "Transient\ndispersion", "Bulk layers"]):
    vial(x, vy, kind=k, seed=1)
    ov.text(x + 4.5, vy + 13.3, t, ha="center", va="top", fontsize=6, linespacing=1.05)
for a, b_, t in [(vx[0] + 10, vx[1] - 1, "demixing"), (vx[1] + 10, vx[2] - 1, "coarsening")]:
    ov.add_patch(FancyArrowPatch((a, vy + 6), (b_, vy + 6), arrowstyle="-|>", mutation_scale=5, lw=0.75, color=GREY))
    ov.text((a + b_) / 2, vy + 5.2, t, ha="center", va="bottom", fontsize=5.5, color=GREY)
ov.text((vx[1] + vx[2] + 9) / 2, vy + 7, "min–h", ha="center", va="top", fontsize=5.5, color=GREY)

hx, hy = vx[1], BY + 30
vial(hx, hy, kind="held")
ov.text(hx + 4.5, hy + 13.3, "Held emulsion", ha="center", va="top", fontsize=6, color=VERM, fontweight="bold")
ov.text(hx + 4.5, hy + 16.0, "weeks; stored, operated", ha="center", va="top", fontsize=5.5, color=VERM)
ov.add_patch(FancyArrowPatch((vx[0] + 4.5, vy + 16.8), (hx - 1, hy + 6), arrowstyle="-|>", mutation_scale=5, lw=0.75,
                             color=VERM, connectionstyle="arc3,rad=0.35"))
ov.add_patch(FancyArrowPatch((vx[2] + 4.5, vy + 16.8), (hx + 10, hy + 6), arrowstyle="-|>", mutation_scale=5, lw=0.75,
                             color=VERM, connectionstyle="arc3,rad=-0.35"))
ov.text(BX, hy + 9.5, "+ CAP in feed:\nnanodroplets", ha="left", va="top", fontsize=5.5, color=VERM, linespacing=1.05)
ov.text(BX + 61, hy + 9.5, "+ CAP and shear:\nµm–mm droplets", ha="right", va="top", fontsize=5.5, color=VERM,
        linespacing=1.05)

# ======================= panel c: CAP design handle
CX, CY = 119, 64
# interface cartoon
iw, top, mid, bot = 22, CY + 1, CY + 13, CY + 27
ov.add_patch(Rectangle((CX, top), iw, mid - top, fc=WATER, ec="none", zorder=1))
ov.add_patch(Rectangle((CX, mid), iw, bot - mid, fc=T_FILL, ec="none", zorder=1))
ov.text(CX + 0.6, top + 0.6, "water", fontsize=5.5, va="top", color="#333333")
ov.text(CX + 0.6, bot - 0.6, "coacervate", fontsize=5.5, va="bottom", color="#1d6b61")
for k, px in enumerate([CX + 5.5, CX + 12, CX + 18.5]):
    t = np.linspace(0, 1, 60)
    ov.plot(px + np.sin(t * 5 * np.pi) * 1.1, mid - 0.4 - t * 9.5, color=BLUE, lw=0.75, zorder=3)
    beads = [(px + 0.2, mid + 1.1), (px - 0.8, mid + 2.4), (px + 0.5, mid + 3.6), (px - 0.4, mid + 5.0), (px + 0.6, mid + 6.3)]
    ov.plot(*zip(*([(px, mid - 0.4)] + beads)), color=VERM, lw=0.5, zorder=3)
    for j, (bx, by) in enumerate(beads):
        dot(ov, bx, by, 0.55, VERM if j % 2 == 0 else "white", z=4, ec=VERM, lw=0.4)
ov.text(CX + iw + 0.8, top + 4.5, "PEG corona", fontsize=5.5, color=BLUE, va="center")
ov.text(CX + iw + 0.8, mid + 3.8, "B/M anchor", fontsize=5.5, color=VERM, va="center")
ov.text(CX, bot + 1.3, "B, phenylboronic acid (filled)\nM, amidoamine (open), random", fontsize=5, va="top",
        color="#333333", linespacing=1.1)


def morph_box(x, y, s, cont):
    fc_bg, fc_d = (T_FILL, WATER) if cont == "C" else (WATER, T_FILL)
    ov.add_patch(Rectangle((x, y), s, s, fc=fc_bg, ec="#555555", lw=0.4, zorder=2))
    for i in range(3):
        for j in range(3):
            ov.add_patch(Circle((x + s * (0.2 + 0.3 * i), y + s * (0.2 + 0.3 * j)), s * 0.1, fc=fc_d, ec=VERM, lw=0.5, zorder=3))


mx0, ms = CX + 35, 9.5
morph_box(mx0, CY + 3, ms, "C")
morph_box(mx0 + 15, CY + 3, ms, "W")
ov.text(mx0 + ms / 2, CY + 2.3, "W/C", ha="center", va="bottom", fontsize=6)
ov.text(mx0 + 15 + ms / 2, CY + 2.3, "C/W", ha="center", va="bottom", fontsize=6)
ax_y = CY + 16.5
ov.add_patch(FancyArrowPatch((mx0 - 0.5, ax_y), (mx0 + 15 + ms + 0.5, ax_y), arrowstyle="-|>", mutation_scale=5, lw=0.75,
                             color=BLACK))
ov.plot([mx0 + ms + 2.75] * 2, [CY + 2, ax_y - 1], ls=(0, (1.5, 1)), lw=0.5, color=GREY)
ov.text(mx0 + 12.25, ax_y + 1.2, "Corona : anchor ratio", ha="center", va="top", fontsize=6)
ov.text(mx0 - 0.5, ax_y + 3.7, "longer anchor", ha="left", va="top", fontsize=5.5, color=VERM)
ov.text(mx0 + 15 + ms + 0.5, ax_y + 3.7, "larger PEG", ha="right", va="top", fontsize=5.5, color=BLUE)
ov.text(mx0 + 12.25, ax_y + 7.3, "The phase hosting the larger block\nbecomes continuous", ha="center", va="top",
        fontsize=5.5, color="#333333", linespacing=1.1)

# panel letters (shared top baseline for a and b)
letter(fig, 1.0, 1.0, "a")
letter(fig, BX - 3, 1.0, "b")
letter(fig, BX - 3, CY - 2, "c")

for ext in ("svg", "pdf"):
    fig.savefig(HERE / f"Fig1.{ext}")
fig.savefig(HERE / "Fig1_preview.png", dpi=300)
# SVG: name the font Illustrator should use
svg = (HERE / "Fig1.svg").read_text()
(HERE / "Fig1.svg").write_text(svg.replace("'Liberation Sans'", "Arial").replace("Liberation Sans", "Arial"))
print("done")
